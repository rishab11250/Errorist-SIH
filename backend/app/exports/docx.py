"""DOCX renderer for an immutable inspection report."""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Inches

from app.exports.pdf import decode_source_image
from app.exports.view_model import InspectionReport


def render_docx(report: InspectionReport) -> bytes:
    """Render a complete, editable compliance report as DOCX bytes."""
    document = Document()
    document.add_heading("LMPC Compliance Report", 0)
    document.add_paragraph(f"Scan ID: {report.scan_id}")
    document.add_paragraph(f"Product: {report.product_name}")
    document.add_paragraph(f"Inspector: {report.owner_display_name}")
    document.add_paragraph(f"Date: {report.created_at.isoformat()}")
    document.add_paragraph(f"Mode: {report.mode}")
    document.add_paragraph(f"Category: {report.category}")
    document.add_paragraph(f"Overall status: {report.overall_status}")
    document.add_paragraph(f"Analysis version: {report.analysis_version}")
    document.add_paragraph(f"Quality summary: {report.quality_summary}")

    document.add_heading("Source label", level=1)
    image = decode_source_image(report.source_image_b64)
    if image is not None:
        try:
            document.add_picture(image, width=Inches(6))
        except Exception:
            document.add_paragraph("Source image could not be embedded.")
    else:
        document.add_paragraph("Source image could not be decoded.")

    document.add_heading("Rule verdicts", level=1)
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headings = (
        "Rule / citation",
        "Status",
        "Evidence",
        "Failure message",
        "Reasoning",
        "Method",
        "Confidence",
        "Review state",
    )
    for cell, heading in zip(table.rows[0].cells, headings, strict=True):
        cell.text = heading
    for verdict in report.verdicts:
        cells = table.add_row().cells
        values = (
            f"{verdict.rule_id}\n{verdict.citation}",
            verdict.status,
            verdict.evidence,
            verdict.failure_message,
            verdict.reasoning,
            verdict.measurement_method,
            f"{verdict.confidence:.0%}",
            verdict.review_state,
        )
        for cell, value in zip(cells, values, strict=True):
            cell.text = value

    document.add_heading("Review history", level=1)
    if report.reviews:
        for review in report.reviews:
            document.add_paragraph(
                f"{review.created_at.isoformat()} — {review.actor_display_name}: "
                f"{review.action}. {review.note}"
            )
    else:
        document.add_paragraph("No review actions recorded.")
    document.add_paragraph(f"Rules applied: {', '.join(report.rules_versions) or 'n/a'}")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
