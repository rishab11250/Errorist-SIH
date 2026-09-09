"""Tests for auditable PDF, DOCX, and filtered CSV exports."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import db, main
from app.auth.sessions import issue_session
from app.db import ReviewAction, Scan, User, VerdictRow
from app.main import app

_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _scan(
    owner_id: int,
    *,
    product: str,
    mode: str = "retail_image",
    status: str = "fail",
    image_b64: str = _PNG,
    rule_id: str = "r6_1_e_mrp",
    verdict_status: str = "fail",
) -> Scan:
    return Scan(
        owner_user_id=owner_id,
        product_name=product,
        mode=mode,
        category="non_food",
        image_b64=image_b64,
        image_meta={"width": 1, "height": 1},
        ocr_payload=[{"text": product}],
        quality_summary={"usable": True, "guidance": "Image quality is sufficient."},
        overall_status=status,
        analysis_version="inspection-v2",
        verdicts=[
            VerdictRow(
                rule_id=rule_id,
                status=verdict_status,
                severity="critical",
                citation="Rule 6(1)(e) of LMPC Rules 2011",
                evidence="MRP inclusive of all taxes",
                evidence_bboxes=[[0.1, 0.1, 0.5, 0.2]],
                failure_message="Required declaration needs attention",
                rule_version="2026-09",
                confidence=0.91,
                reasoning="The required phrase was evaluated from OCR evidence.",
                measurement_method="relative_text_height",
                review_state="confirmed",
            )
        ],
    )


@pytest.fixture
def export_context(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DB_PATH", tmp_path / "exports.db")
    with TestClient(app) as client:
        assert db.SessionLocal is not None
        with db.SessionLocal() as session:
            inspector = User(
                username_normalized="inspector",
                display_name="Inspector One",
                password_hash="unused",
                role="inspector",
            )
            other = User(
                username_normalized="other",
                display_name="Other Inspector",
                password_hash="unused",
                role="inspector",
            )
            admin = User(
                username_normalized="admin",
                display_name="Administrator",
                password_hash="unused",
                role="admin",
            )
            session.add_all([inspector, other, admin])
            session.flush()
            tokens = {
                user.username_normalized: issue_session(
                    session,
                    user,
                    now=datetime.now(UTC),
                    ttl=timedelta(hours=8),
                ).token
                for user in (inspector, other, admin)
            }
            reviewed = _scan(inspector.id, product="चाय Premium", mode="ecommerce_listing")
            formula = _scan(inspector.id, product='=HYPERLINK("https://bad")', status="pass")
            invalid_image = _scan(inspector.id, product="Invalid image", image_b64="not-base64")
            mixed = _scan(inspector.id, product="Mixed result", status="mixed")
            manual = _scan(inspector.id, product="Manual review", status="manual_review")
            other_scan = _scan(other.id, product="Other owner")
            session.add_all([reviewed, formula, invalid_image, mixed, manual, other_scan])
            session.flush()
            session.add(
                ReviewAction(
                    scan_id=reviewed.id,
                    verdict_id=reviewed.verdicts[0].id,
                    actor_user_id=inspector.id,
                    action="confirmed",
                    note="Verified against the physical pack.",
                )
            )
            session.commit()
            result = {
                "client": client,
                "tokens": tokens,
                "reviewed_id": reviewed.id,
                "formula_id": formula.id,
                "invalid_image_id": invalid_image.id,
                "other_id": other_scan.id,
                "inspector_id": inspector.id,
            }
        yield result


def _as(context, username: str) -> TestClient:
    client = context["client"]
    client.cookies.clear()
    client.cookies.set("lmpc_session", context["tokens"][username])
    return client


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def test_pdf_and_docx_contain_same_core_report_values(export_context) -> None:
    client = _as(export_context, "admin")
    scan_id = export_context["reviewed_id"]
    pdf = client.get(f"/api/exports/scans/{scan_id}.pdf")
    docx = client.get(f"/api/exports/scans/{scan_id}.docx")
    assert pdf.status_code == docx.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-type"] == "application/pdf"
    assert docx.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    text = _docx_text(docx.content)
    for value in (str(scan_id), "fail", "Rule 6", "चाय Premium", "Verified against"):
        assert value in text
    assert pdf.headers["content-disposition"] == f'attachment; filename="lmpc-scan-{scan_id}.pdf"'
    assert docx.headers["content-disposition"] == (
        f'attachment; filename="lmpc-scan-{scan_id}.docx"'
    )


def test_csv_neutralizes_formula_cells_and_preserves_unicode(export_context) -> None:
    response = _as(export_context, "admin").get("/api/exports/scans.csv")
    assert response.status_code == 200
    assert "'=HYPERLINK" in response.text
    assert "चाय Premium" in response.text
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert len(rows) == 6
    assert response.headers["x-lmpc-result-count"] == "6"
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"].startswith('attachment; filename="lmpc-scans-')


def test_csv_uses_authorized_active_filters_and_digest(export_context) -> None:
    client = _as(export_context, "inspector")
    response = client.get(
        "/api/exports/scans.csv?mode=ecommerce_listing&rule_id=r6_1_e_mrp&verdict_status=fail"
    )
    assert response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert [row["scan_id"] for row in rows] == [str(export_context["reviewed_id"])]
    active = {
        "mode": "ecommerce_listing",
        "rule_id": "r6_1_e_mrp",
        "verdict_status": "fail",
    }
    expected = hashlib.sha256(
        json.dumps(active, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    assert response.headers["x-lmpc-filter-digest"] == expected
    assert response.headers["x-lmpc-result-count"] == "1"


@pytest.mark.parametrize("status", ["pass", "fail", "mixed", "manual_review"])
def test_csv_supports_every_overall_status(export_context, status: str) -> None:
    response = _as(export_context, "admin").get(f"/api/exports/scans.csv?overall_status={status}")
    assert response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert rows
    assert {row["overall_status"] for row in rows} == {status}


def test_inspector_cannot_export_another_scan_or_filter_by_owner(export_context) -> None:
    client = _as(export_context, "inspector")
    assert client.get(f"/api/exports/scans/{export_context['other_id']}.pdf").status_code == 404
    assert client.get(f"/api/exports/scans/{export_context['other_id']}.docx").status_code == 404
    response = client.get(f"/api/exports/scans.csv?owner_id={export_context['inspector_id']}")
    assert response.status_code == 403


def test_exports_handle_missing_scans_and_invalid_source_image(export_context) -> None:
    client = _as(export_context, "admin")
    assert client.get("/api/exports/scans/99999.pdf").status_code == 404
    assert client.get("/api/exports/scans/99999.docx").status_code == 404
    scan_id = export_context["invalid_image_id"]
    assert client.get(f"/api/exports/scans/{scan_id}.pdf").content.startswith(b"%PDF")
    docx = client.get(f"/api/exports/scans/{scan_id}.docx")
    assert docx.status_code == 200
    assert "Invalid image" in _docx_text(docx.content)


def test_legacy_report_route_uses_export_service(export_context) -> None:
    scan_id = export_context["reviewed_id"]
    response = _as(export_context, "admin").get(f"/api/report/{scan_id}")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-disposition"] == (
        f'attachment; filename="lmpc-scan-{scan_id}.pdf"'
    )
